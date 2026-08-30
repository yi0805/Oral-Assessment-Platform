import asyncio
import logging
from datetime import datetime, timedelta, timezone
from io import BytesIO
from uuid import UUID

from sqlalchemy import text as sa_text
from sqlalchemy.orm import Session

from app.core.config import settings
from app.core.database import SessionLocal
from app.models import Material, MaterialChunk

from app.services import s3_client
from app.services.embedding_service import embed_batch

from app.utils.pdf_extractor import extract_text_from_bytes
from app.utils.text_chunker import chunk_text

# 3-stage pipeline (after upload): EXTRACT → CHUNK → EMBED

logger = logging.getLogger(__name__)


def mark_material_processing(material: Material, *, now: datetime | None = None) -> None:
    """Start a processing attempt and record the time used for stale recovery."""
    material.processing_status = "processing"
    material.processing_started_at = now or datetime.now(timezone.utc)


def is_material_processing_stale(
    material: Material, *, now: datetime | None = None
) -> bool:
    """Return true only for an old, non-terminal processing attempt."""
    if material.processing_status != "processing" or material.processing_started_at is None:
        return False

    started_at = material.processing_started_at
    if started_at.tzinfo is None:
        started_at = started_at.replace(tzinfo=timezone.utc)
    current_time = now or datetime.now(timezone.utc)
    threshold = timedelta(seconds=settings.material_processing_stale_seconds)
    return started_at <= current_time - threshold


def can_retry_material_processing(material: Material, *, now: datetime | None = None) -> bool:
    return material.processing_status == "failed" or is_material_processing_stale(
        material, now=now
    )


def _get_db() -> Session:

    return SessionLocal()


def _run_async(coro):

    try:
        loop = asyncio.get_running_loop()

    except RuntimeError:
        return asyncio.run(coro)

    logger.warning("Running event loop detected in pipeline thread — using threadsafe dispatch")
    future = asyncio.run_coroutine_threadsafe(coro, loop)
    return future.result(timeout=180)


def _file_type_from_filename(filename: str) -> str:

    if "." in filename:
        return filename.rsplit(".", 1)[-1].lower()
    return ""

def _extract_docx(file_bytes: bytes) -> tuple[str, int]:
    try:
        import docx  
    except ImportError:
        raise RuntimeError("python-docx not installed. Run: pip install python-docx")

    doc = docx.Document(BytesIO(file_bytes))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(c.text.strip() for c in row.cells if c.text.strip())
            if row_text:
                paragraphs.append(row_text)
    full_text = "\n\n".join(paragraphs)
    word_count = len(full_text.split())
    return full_text, max(1, round(word_count / 250))


def _extract_pptx(file_bytes: bytes) -> tuple[str, int]:
    try:
        from pptx import Presentation
    except ImportError:
        raise RuntimeError("python-pptx not installed. Run: pip install python-pptx")

    prs = Presentation(BytesIO(file_bytes))
    slides_text: list[str] = []
    for slide_no, slide in enumerate(prs.slides, start=1):
        parts = [
            para.text.strip()
            for shape in slide.shapes
            if shape.has_text_frame
            for para in shape.text_frame.paragraphs
            if para.text.strip()
        ]
        if parts:
            slides_text.append(f"[Slide {slide_no}]\n" + "\n".join(parts))
    return "\n\n".join(slides_text), len(prs.slides)



# ---------------------------------------------------------------------------
# Public entry point
# ---------------------------------------------------------------------------


def run_pipeline(material_id: UUID) -> None:
    """Run the existing pipeline and publish one truthful terminal status."""
    logger.info("Pipeline start for material %s", material_id)
    _set_material_status(material_id, "processing")

    try:
        extracted_text = _stage_extract(material_id)
        if extracted_text is None:
            _set_material_status(material_id, "failed")
            return

        if not _stage_chunk(material_id, extracted_text):
            _set_material_status(material_id, "failed")
            return

        if not _stage_embed(material_id):
            _set_material_status(material_id, "failed")
            return

        _set_material_status(material_id, "ready")
        logger.info("Pipeline finished for material %s", material_id)
    except Exception:
        logger.exception("Unexpected pipeline failure for material %s", material_id)
        _set_material_status(material_id, "failed")


def _set_material_status(material_id: UUID, processing_status: str) -> None:
    db = _get_db()
    try:
        material = db.query(Material).filter(Material.id == material_id).first()
        if material:
            if processing_status == "processing":
                mark_material_processing(material)
            else:
                material.processing_status = processing_status
                material.processing_started_at = None
            db.commit()
    except Exception:
        db.rollback()
        logger.exception(
            "Could not set processing status for material %s", material_id
        )
    finally:
        db.close()


def _stage_extract(material_id: UUID) -> str | None:

    db = _get_db()

    try:
        material = db.query(Material).filter(Material.id == material_id).first()

        if not material:
            logger.error("Material %s not found during extraction", material_id)
            return None

        try:
            file_bytes = s3_client.download_file(material.storage_key)

        except FileNotFoundError:
            logger.error(
                "Material %s: file not found in storage: %s",
                material_id, material.storage_key,
            )
            return None

        except Exception as exc:
            logger.error("Material %s: storage download error: %s", material_id, exc)
            return None

        file_type = _file_type_from_filename(material.filename)

        try:
            if file_type == "pdf":
                result = extract_text_from_bytes(file_bytes)
                extracted_text = result.text

            elif file_type == "txt":
                extracted_text = file_bytes.decode("utf-8", errors="replace")

            elif file_type == "docx":
                extracted_text, _ = _extract_docx(file_bytes)

            elif file_type == "pptx":
                extracted_text, _ = _extract_pptx(file_bytes)

            else:
                logger.error("Material %s: unsupported file type: %s", material_id, file_type)
                return None

        except Exception as exc:  
            logger.error(
                "Material %s: extraction error (%s): %s: %s",
                material_id, file_type, type(exc).__name__, exc,
            )
            return None

        if not extracted_text or not extracted_text.strip():
            logger.error("Material %s: extraction returned empty text", material_id)
            return None

        logger.info("Material %s: text extracted (%d chars)", material_id, len(extracted_text))

        return extracted_text

    except Exception as exc:  
        logger.exception("Unexpected error in _stage_extract for %s", material_id)
        return None
    
    finally:
        try:
            db.close()

        except Exception:  
            pass


def _stage_chunk(material_id: UUID, extracted_text: str) -> bool:

    db = _get_db()

    try:
        chunks = chunk_text(extracted_text, max_tokens=500, overlap_tokens=50)
        if not chunks:
            logger.error("Material %s: chunking produced zero chunks", material_id)
            return False

        db.execute(
            sa_text("DELETE FROM material_chunks WHERE material_id = :mid"),
            {"mid": str(material_id)},
        )

        for chunk in chunks:
            db.add(MaterialChunk(
                material_id=material_id,
                chunk_index=chunk.chunk_index,
                chunk_text=chunk.chunk_text,
            ))

        db.commit()
        logger.info("Material %s: %d chunks saved", material_id, len(chunks))

    except Exception as exc: 
        logger.exception("Unexpected error in _stage_chunk for %s", material_id)
        return False
    
    finally:
        try:
            db.close()

        except Exception: 
            pass

    return True



def _stage_embed(material_id: UUID) -> bool:

    db = _get_db()

    try:
        chunks = (
            db.query(MaterialChunk)
            .filter(MaterialChunk.material_id == material_id)
            .order_by(MaterialChunk.chunk_index.asc())
            .all()
        )

        if not chunks:
            logger.warning("Material %s: no chunks to embed", material_id)
            return False

        texts = [c.chunk_text for c in chunks]
        try:
            embeddings = _run_async(embed_batch(texts))

            if len(embeddings) != len(chunks):
                raise ValueError(
                    f"Embedding count mismatch: got {len(embeddings)}, "
                    f"expected {len(chunks)}"
                )
            
        except Exception as exc:
            logger.warning("Material %s: embedding failed: %s", material_id, exc)
            return False

        if any(len(vec) != 768 or all(value == 0.0 for value in vec) for vec in embeddings):
            logger.warning("Material %s: embedding response is unusable", material_id)
            return False

        for chunk, vec in zip(chunks, embeddings):
            chunk.embedding = vec

        db.commit()

        logger.info(
            "Material %s: pipeline complete — %d chunks embedded (768-dim Gemini)",
            material_id, len(chunks),
        )
        return True

    except Exception as exc:  
        logger.exception("Unexpected error in _stage_embed for %s", material_id)
        return False
    finally:
        try:
            db.close()
        except Exception: 
            pass
